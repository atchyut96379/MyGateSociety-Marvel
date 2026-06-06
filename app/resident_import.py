from dataclasses import dataclass
from io import BytesIO
import re
from typing import Any

from openpyxl import Workbook, load_workbook


@dataclass(frozen=True)
class ImportedResident:
    row_number: int
    tower: str
    flat_number: str
    name: str
    phone: str
    email: str | None
    role: str
    property_owner_name: str | None


HEADER_ALIASES = {
    "flat": {
        "flat",
        "flatno",
        "flatnumber",
        "flatnofirst",
        "unit",
        "unitnumber",
        "apartment",
        "apartmentnumber",
    },
    "name": {
        "name",
        "resident",
        "residentname",
        "membername",
        "tenantname",
        "ownername",
    },
    "phone": {
        "phone",
        "mobile",
        "mobilenumber",
        "loginmobile",
        "contact",
        "contactnumber",
        "residentmobile",
    },
    "owner_phone": {
        "ownercontactnumber",
        "ownermobile",
    },
    "tenant_phone": {
        "tenantcontactnumber",
        "tenantmobile",
    },
    "email": {
        "email",
        "notificationemail",
        "emailaddress",
    },
    "role": {
        "type",
        "residenttype",
        "membertype",
    },
    "property_owner_name": {
        "propertyowner",
        "propertyownername",
        "owner",
    },
}


def _key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _phone(value: Any) -> str:
    text = _clean(value)
    if not text:
        return ""
    if text.endswith(".0"):
        text = text[:-2]
    text = re.sub(r"[^\d+]", "", text)
    if text.startswith("91") and len(text) == 12:
        text = f"+{text}"
    return text


def _role(value: Any) -> str:
    text = _clean(value).lower()
    if "owner" in text:
        return "owner"
    if "tenant" in text:
        return "tenant"
    if "family" in text:
        return "family"
    return "member"


def _split_flat(value: Any, default_tower: str) -> tuple[str, str]:
    flat = _clean(value)
    if not flat:
        return default_tower, ""

    for separator in ("\\", "/", "-", " "):
        if separator in flat:
            first, rest = flat.split(separator, 1)
            if any(ch.isalpha() for ch in first) and rest.strip():
                return first.strip().upper(), rest.strip()
    return default_tower, flat


def _map_headers(row: tuple[Any, ...]) -> dict[str, int]:
    mapped: dict[str, int] = {}
    for index, value in enumerate(row):
        header = _key(value)
        if not header:
            continue
        for field, aliases in HEADER_ALIASES.items():
            if header in aliases and field not in mapped:
                mapped[field] = index
    return mapped


def _best_sheet(workbook) -> tuple[Any, int, dict[str, int]] | None:
    best: tuple[Any, int, dict[str, int]] | None = None
    best_score = 0
    for sheet in workbook.worksheets:
        for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
            mapped = _map_headers(row)
            score = len(set(mapped) & {"flat", "name", "phone", "role"})
            if score > best_score and {"flat", "name"} <= set(mapped):
                best = (sheet, row_number, mapped)
                best_score = score
    return best


def parse_resident_workbook(content: bytes, default_tower: str = "A") -> tuple[list[ImportedResident], list[str]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    selection = _best_sheet(workbook)
    if selection is None:
        return [], ["Could not find a sheet with Flat Number and Resident Name columns."]

    sheet, header_row, mapped = selection
    residents: list[ImportedResident] = []
    errors: list[str] = []
    seen: set[tuple[str, str, str, str]] = set()

    for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if not any(_clean(value) for value in row):
            continue

        def value(field: str) -> Any:
            index = mapped.get(field)
            return row[index] if index is not None and index < len(row) else None

        tower, flat_number = _split_flat(value("flat"), default_tower)
        name = _clean(value("name"))
        email = _clean(value("email")) or None
        role = _role(value("role"))
        if role == "tenant":
            phone = _phone(value("tenant_phone") or value("phone") or value("owner_phone"))
        else:
            phone = _phone(value("owner_phone") or value("phone") or value("tenant_phone"))
        property_owner_name = _clean(value("property_owner_name")) or None

        if not flat_number:
            errors.append(f"Row {row_number}: Flat Number is required.")
            continue
        if not name:
            errors.append(f"Row {row_number}: Resident Name is required.")
            continue
        if not phone:
            errors.append(f"Row {row_number}: Login mobile/phone is required.")
            continue

        duplicate_key = (tower.lower(), flat_number.lower(), name.lower(), role)
        if duplicate_key in seen:
            errors.append(f"Row {row_number}: Duplicate resident {name} for flat {flat_number} in file.")
            continue
        seen.add(duplicate_key)

        residents.append(
            ImportedResident(
                row_number=row_number,
                tower=tower,
                flat_number=flat_number,
                name=name,
                phone=phone,
                email=email,
                role=role,
                property_owner_name=property_owner_name,
            )
        )

    return residents, errors


def build_resident_import_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Residents"
    sheet.append(["FlatNumber", "ResidentName", "Phone", "NotificationEmail", "ResidentType", "PropertyOwnerName"])
    sheet.append(["101", "Asha Rao", "+919999000001", "asha@example.com", "Owner", ""])
    sheet.append(["102", "Rajendra Kumar T", "+919999000004", "", "Tenant", "Asha Rao"])

    society_sheet = workbook.create_sheet("SocietyFormat")
    society_sheet.append(["SNO", "Flat number", "Name", "Resident Type", "Owner Contact Number", "Tenant Contact Number"])
    society_sheet.append([1, "101", "Asha Rao", "Owner", "+919999000001", ""])
    society_sheet.append([2, "102", "Rajendra Kumar T", "Tenant", "", "+919999000004"])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
